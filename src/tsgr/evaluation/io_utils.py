"""Small deterministic CSV/XLSX helpers used by experiment reports."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Iterable, Sequence


def read_csv_rows(path: str | Path) -> list[dict[str, str]]:
    candidate = Path(path)
    if not candidate.is_file():
        return []
    with candidate.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def write_csv_rows(path: str | Path, rows: Iterable[dict[str, Any]], fields: Sequence[str]) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(fields), extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_workbook(path: str | Path, sheets: dict[str, tuple[Sequence[str], Sequence[dict[str, Any]]]]) -> None:
    """Write compact research tables to XLSX; frame-level data intentionally stays in CSV."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError as error:  # pragma: no cover - environment guard
        raise RuntimeError(
            "XLSX export requires openpyxl. Install the pinned requirements.txt environment."
        ) from error

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    for sheet_name, (fields, rows) in sheets.items():
        worksheet = workbook.create_sheet(title=str(sheet_name)[:31])
        worksheet.append(list(fields))
        for row in rows:
            worksheet.append([row.get(field, "") for field in fields])
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        worksheet.freeze_panes = "A2"
        for column_index, field in enumerate(fields, start=1):
            width = max(10, min(34, len(str(field)) + 2))
            for row_index in range(2, min(worksheet.max_row, 250) + 1):
                value = worksheet.cell(row=row_index, column=column_index).value
                if value is not None:
                    width = min(40, max(width, len(str(value)) + 2))
            worksheet.column_dimensions[get_column_letter(column_index)].width = width
    workbook.save(target)
